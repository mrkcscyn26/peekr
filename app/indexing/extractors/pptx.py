"""C5 PPTX extractor. Implements F2, F5. Covers FR-4, FR-5. Built in T-03.

One segment per slide (titles, text boxes, tables, grouped shapes, speaker notes),
location 'slide N' (Section 8.10).
"""

from __future__ import annotations

from typing import Iterable, Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.indexing.extractors import Segment, check_not_encrypted_ooxml


def _shape_texts(shapes: Iterable) -> Iterator[str]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_texts(shape.shapes)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    yield " | ".join(cells)
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield shape.text_frame.text


def extract(path: str) -> list[Segment]:
    check_not_encrypted_ooxml(path)
    with open(path, "rb") as fh:
        prs = Presentation(fh)
    segments = []
    for i, slide in enumerate(prs.slides, 1):
        parts = list(_shape_texts(slide.shapes))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            parts.append(slide.notes_slide.notes_text_frame.text)
        segments.append(Segment("\n".join(p for p in parts if p.strip()), f"slide {i}"))
    return segments
